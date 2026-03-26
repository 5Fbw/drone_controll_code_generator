import json
import os
import sys
import argparse
import openpyxl


def parse_json_string(value):
    """
    尝试将字符串解析为JSON对象（列表或字典）。
    如果解析失败或不是字符串，则返回原值。
    """
    if not isinstance(value, str):
        return value

    # 去除首尾空格
    value = value.strip()

    # 快速判断是否像 JSON 结构
    if (value.startswith('[') and value.endswith(']')) or \
            (value.startswith('{') and value.endswith('}')):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            # 解析失败，返回原字符串
            return value
    return value


def excel_to_json(input_file: str, output_file: str):
    """
    将Excel文件转换为JSON文件，并尝试自动还原列表/字典类型字段
    """
    print(f"读取输入文件: {input_file}")

    if not os.path.exists(input_file):
        print(f"错误: 输入文件不存在: {input_file}")
        return

    try:
        wb = openpyxl.load_workbook(input_file)
        ws = wb.active
    except Exception as e:
        print(f"错误: 无法读取Excel文件: {e}")
        return

    # 读取表头
    headers = []
    for cell in ws[1]:
        if cell.value is not None:
            headers.append(str(cell.value))  # 确保表头是字符串
        else:
            headers.append(f"Unnamed_{cell.column}")

    data = []

    # 从第二行开始遍历
    for row_values in ws.iter_rows(min_row=2, values_only=True):
        row_data = {}

        # 跳过全空行
        if all(value is None for value in row_values):
            continue

        for col_idx, header in enumerate(headers):
            # 获取值，处理列数不一致的情况
            value = row_values[col_idx] if col_idx < len(row_values) else None

            # 处理空值
            if value is None:
                value = ""

            # 核心：尝试将字符串解析为列表或字典
            value = parse_json_string(value)

            row_data[header] = value

        data.append(row_data)

    print(f"共读取 {len(data)} 条记录")

    output_dir = os.path.dirname(output_file)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)

    try:
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=4)
        print(f"结果已保存到: {output_file}")
    except Exception as e:
        print(f"保存文件时出错: {e}")


def main_cli():
    parser = argparse.ArgumentParser(description='Excel转JSON工具（支持类型还原）')
    parser.add_argument('--input', type=str, default='output_qwen3_32b_all_0318_agent.xlsx',
                        help='输入的Excel文件路径')
    parser.add_argument('--output', type=str, default='output_qwen3_32b_all_0318_1_agent.json',
                        help='输出的JSON文件路径')
    args = parser.parse_args()

    excel_to_json(args.input, args.output)


if __name__ == '__main__':
    main_cli()
