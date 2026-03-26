import json
import os
import sys
import argparse
import openpyxl


def json_to_excel(input_file: str, output_file: str):
    """
    将JSON文件转换为Excel文件

    Args:
        input_file: 输入JSON文件路径
        output_file: 输出Excel文件路径
    """
    print(f"读取输入文件: {input_file}")
    with open(input_file, 'r', encoding='utf-8') as f:
        data = json.load(f)

    print(f"共读取 {len(data)} 条记录")

    # 创建Excel工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "无人机控制代码"

    # 定义所有字段
    all_fields = ['case_id', 'instruction', 'origin_pos', 'exp_destination',
                  'exp_path', 'code', 'path_file', 'multistep', 'actual_path', 'model',
                  'token', 'prompt_tokens', 'completion_tokens']

    # 写入表头
    for col, header in enumerate(all_fields, 1):
        ws.cell(row=1, column=col, value=header)

    # 写入数据
    for row, item in enumerate(data, 2):
        for col, field in enumerate(all_fields, 1):
            value = item.get(field, '')
            # 处理列表类型的数据
            if isinstance(value, list):
                value = str(value)
            
            cell = ws.cell(row=row, column=col, value=value)
            # 代码列和其他可能较长的文本列设置自动换行
            if field in ['code', 'instruction', 'origin_pos', 'exp_destination', 'exp_path', 'actual_path']:
                cell.alignment = openpyxl.styles.Alignment(wrapText=True)

    # 调整列宽
    column_widths = {
        'case_id': 10,
        'instruction': 60,
        'origin_pos': 20,
        'exp_destination': 20,
        'exp_path': 40,
        'code': 80,
        'path_file': 30,
        'multistep': 10,
        'actual_path': 40,
        'model': 20,
        'token': 12,
        'prompt_tokens': 15,
        'completion_tokens': 18
    }

    for col, field in enumerate(all_fields, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = column_widths.get(field, 20)

    # 确保输出目录存在
    output_dir = os.path.dirname(output_file)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)

    # 保存Excel文件
    wb.save(output_file)
    print(f"结果已保存到: {output_file}")


def main_cli():
    parser = argparse.ArgumentParser(description='JSON转Excel工具')

    # 修改此处：删除 required=True，添加 default 参数
    parser.add_argument('--input', type=str, default='output_qwen3_32b_all_0318_1_agent.json',
                        help='输入的JSON文件路径')

    parser.add_argument('--output', type=str, default='output_qwen3_32b_all_0318_1_agent.xlsx',
                        help='输出的Excel文件路径')

    args = parser.parse_args()

    input_path = args.input
    output_path = args.output

    if not os.path.exists(input_path):
        print(f"错误: 输入文件不存在: {input_path}")
        sys.exit(1)

    json_to_excel(input_path, output_path)


if __name__ == '__main__':
    main_cli()
